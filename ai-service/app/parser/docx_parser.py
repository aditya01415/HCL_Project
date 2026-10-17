from docx import Document
from pathlib import Path

def extract_text_from_docx(file_path: str | Path) -> str:
    """
    Extract text content from a DOCX document using python-docx.
    """
    file_path = str(file_path)
    text = ""
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text += paragraph.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"[Error] Failed to extract text from DOCX '{file_path}': {e}")
    return text.strip()
